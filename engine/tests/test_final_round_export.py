"""Final-round export: ESCAP's OUTPUT_TEMPLATE_FINAL_ROUND in Excel, CSV and JSON."""
from __future__ import annotations

import csv
import json
from datetime import datetime, timezone

import openpyxl

from packages.core.finalization import apply_decision
from packages.core.schemas import MappedFinding, ReviewDecision
from packages.export.final_round import (PILLAR_COLUMN, TEMPLATE_COLUMNS, TEMPLATE_PATH, economy_name,
                                         export_final_round, indicator_code, short_rationale,
                                         source_language)


def _approved(economy: str, indicator: str, article: str, snippet: str, **extra) -> MappedFinding:
    finding = MappedFinding.model_validate({
        "Economy": economy, "Law Name": "Personal Data Protection Act 2012",
        "Indicator ID": indicator, "Article / Section": article, "Discovery Tag": "NEW",
        "Location Reference": "page 3", "Verbatim Snippet": snippet,
        "Mapping Rationale": extra.pop("rationale", "This s. 26 permits transfer subject to conditions."),
        "Source URL": "https://sso.agc.gov.sg/Act/PDPA2012", "Confidence": 0.91, "Status": "in_force", **extra,
    })
    return apply_decision(finding, ReviewDecision(
        decision="approved", reviewer_name="Reviewer One", reviewer_role="admin",
        reviewed_at=datetime(2026, 9, 30, tzinfo=timezone.utc), citation_checked=True,
        mapping_checked=True, status_checked=True, correction_note="Genuine conditional-transfer rule.",
        citation_reviewer_name="Reviewer One", mapping_reviewer_name="Reviewer Two",
        status_reviewer_name="Reviewer One"))


def test_template_conventions():
    assert indicator_code("P6-I4") == "6.4"
    assert indicator_code("P12-I10") == "12.10"   # stays distinct from 12.1
    assert indicator_code("6.4") == "6.4"
    assert economy_name("Lao PDR") == "Lao People's Democratic Republic"
    assert source_language("มาตรา ๒๖ ผู้ควบคุมข้อมูลส่วนบุคคล", "Thailand") == ("Thai", "Thai")
    assert source_language("Хувийн мэдээллийг хамгаалах тухай хууль", "Mongolia") == ("Mongolian", "Mongolian")
    assert source_language("Оператор обязан обеспечить запись персональных данных", "Russian Federation")[0] == "Russian"
    assert source_language("An organisation shall not transfer any personal data to a country", "Singapore") == ("English", "English")
    assert source_language("Pengendali Data Pribadi wajib menunjuk pejabat", "Indonesia")[0] == "Bahasa Indonesia"
    assert source_language("O fornecedor nacional beneficia de preferência", "Timor-Leste") == ("Other", "Portuguese")
    long = "word " * 120
    assert len(short_rationale(long)) <= 300 and short_rationale(long).endswith("…")
    assert short_rationale("Short rationale.") == "Short rationale."


def _pending(economy: str, indicator: str, article: str, snippet: str) -> MappedFinding:
    return MappedFinding.model_validate({
        "Economy": economy, "Law Name": "Personal Data Protection Act 2012", "Indicator ID": indicator,
        "Article / Section": article, "Discovery Tag": "NEW", "Location Reference": "page 5",
        "Verbatim Snippet": snippet, "Mapping Rationale": "Localisation duty.",
        "Source URL": "https://sso.agc.gov.sg/Act/PDPA2012", "Confidence": 0.74, "Status": "in_force"})


SCORE_CELLS = [
    {"economy": "Singapore", "indicator": "P6-I4", "state": "approved", "effective": "1.0", "deterministic": "1.0",
     "deterministic_reason": "conditional flow regime evidenced", "reasoning": "Conditional flow regime.",
     "reviewer_name": "Reviewer Two", "reviewed_at": "2026-09-30T01:00:00+00:00", "judge_scores": "strict: 1.0",
     "agreement_alpha": "1.0", "master_gold": "1.0"},
    {"economy": "Singapore", "indicator": "P7-I2", "state": "pending", "effective": None, "deterministic": "0.5",
     "deterministic_reason": "partial localisation", "reasoning": "", "reviewer_name": "", "master_gold": None},
]


def test_export_writes_template_xlsx_csv_and_json(tmp_path):
    entries = [
        (_approved("Singapore", "P6-I4", "s. 26(1)", "An organisation shall not transfer any personal data "
                   "to a country outside Singapore except in accordance with requirements prescribed."), None),
        (_approved("Lao PDR", "P7-I1", "Art. 9", "ຂໍ້ມູນສ່ວນບຸກຄົນ ຕ້ອງໄດ້ຮັບການປົກປ້ອງ ແລະ ຮັກສາໄວ້ ຢ່າງປອດໄພ",
                   rationale="x " * 200), None),
        (_pending("Singapore", "P7-I2", "s. 24", "An organisation shall protect personal data in its possession."),
         {"decision": None, "citation_reviewer_name": "Reviewer One"}),
    ]
    manifest = export_final_round(entries, tmp_path, mode="hybrid", score_cells=SCORE_CELLS,
                                  source={"id": "snap-1", "source_hash": "c" * 64},
                                  excluded={"rejected": 4, "absence": 2})
    assert manifest["rows"] == 3 and manifest["approved_rows"] == 2 and manifest["pending_rows"] == 1
    assert manifest["excluded_rejected"] == 4 and manifest["excluded_no_evidence_placeholders"] == 2
    assert manifest["indicator_scores"] == 2 and manifest["indicator_scores_decided"] == 1
    assert manifest["economies"] == ["Lao People's Democratic Republic", "Singapore"]
    assert {f"{fmt}_sha256" for fmt in ("xlsx", "csv", "json")} <= set(manifest)

    wb = openpyxl.load_workbook(tmp_path / "OUTPUT_FINAL_ROUND.xlsx")
    ws = wb["Output Data"]
    template = openpyxl.load_workbook(TEMPLATE_PATH)["Output Data"]
    assert tuple(ws.cell(4, c).value for c in range(1, 16)) == TEMPLATE_COLUMNS
    assert [ws.cell(4, c).value for c in range(1, 16)] == [template.cell(4, c).value for c in range(1, 16)]
    # Example rows removed: data starts at row 7, sorted by economy, pillar, indicator.
    assert ws["A7"].value == "Lao People's Democratic Republic" and ws["A8"].value == "Singapore"
    assert ws["A10"].value is None
    assert ws["E8"].value == "6.4" and ws["E8"].number_format == "@" and ws["E9"].value == "7.2"
    assert ws["O8"].value.startswith('=IF($E8="","",')
    assert ws["N7"].value == "Lao" and len(ws["J7"].value) <= 300
    assert ws.freeze_panes == "A7"
    assert "2 approved by named reviewers, 1 awaiting review" in ws["A6"].value
    extra = {r: {ws.cell(4, c).value: ws.cell(r, c).value for c in range(16, ws.max_column + 1)} for r in (8, 9)}
    assert extra[8]["Mapping Reviewer"] == "Reviewer Two" and extra[8]["Review Decision"] == "approved"
    assert extra[8]["Indicator Score"] == 1 and extra[8]["Score Status"] == "reviewer-approved"
    assert extra[8]["Legal Status"] == "In force"
    # Empty and internal columns are not exported.
    for dropped in ("Verbatim Snippet (English)", "OCR Confidence", "Review Date", "Mapping Rationale (full)",
                    "Finding Key", "Review Subject Hash"):
        assert dropped not in extra[8]
    assert extra[9]["Review Decision"] == "pending" and extra[9]["Citation Reviewer"] == "Reviewer One"
    assert extra[9]["Indicator Score"] == 0.5 and extra[9]["Score Status"].startswith("engine proposal")
    assert extra[9]["Score Reasoning"] == "partial localisation"
    coverage = wb["Coverage Matrix"]
    rows = {coverage.cell(r, 1).value: r for r in range(4, coverage.max_row + 1)}
    assert "Lao People's Democratic Republic" in rows and "Lao PDR" not in rows
    assert "$O$7:$O$9" in coverage.cell(rows["Singapore"], 7).value
    scores = wb["Indicator Scores"]
    header = [scores.cell(4, c).value for c in range(1, scores.max_column + 1)]
    assert "Review Date" not in header
    first = dict(zip(header, [scores.cell(5, c).value for c in range(1, scores.max_column + 1)]))
    assert first["Indicator ID"] == "6.4" and first["RDTII Reference Score"] == 1 and first["Evidence Rows Exported"] == 1
    assert wb.sheetnames == ["Output Data", "Coverage Matrix", "Indicator Scores"]

    with (tmp_path / "OUTPUT_FINAL_ROUND.csv").open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        assert tuple(reader.fieldnames[:15]) == TEMPLATE_COLUMNS
        assert "Finding Key" not in reader.fieldnames and "Review Date" not in reader.fieldnames
        rows_csv = list(reader)
    assert rows_csv[1]["Indicator ID"] == "6.4" and rows_csv[1][PILLAR_COLUMN] == "6"

    payload = json.loads((tmp_path / "OUTPUT_FINAL_ROUND.json").read_text())
    assert payload["template_columns"] == list(TEMPLATE_COLUMNS)
    assert payload["rows"][1][PILLAR_COLUMN] == 6
    assert payload["manifest"]["source_snapshot"] == "snap-1"
    assert [s["Score Status"] for s in payload["indicator_scores"]] == ["reviewer-approved",
                                                                        "engine proposal (awaiting review)"]


def test_export_script_writes_one_requested_format(tmp_path):
    import subprocess
    import sys
    from pathlib import Path

    finding = _pending("Singapore", "P7-I2", "s. 24", "An organisation shall protect personal data.")
    payload = {"mode": "local", "rows": [{"finding": finding.model_dump(by_alias=True, mode="json"),
                                          "review": {"decision": None}}], "scores": SCORE_CELLS}
    script = Path(__file__).resolve().parents[1] / "scripts" / "export_final_round.py"
    result = subprocess.run([sys.executable, str(script), "--out", str(tmp_path), "--type", "csv"],
                            input=json.dumps(payload), text=True, capture_output=True, check=True)
    manifest = json.loads(result.stdout.strip().splitlines()[-1])
    assert manifest["rows"] == 1 and manifest["engine_mode"].startswith("Local") and "csv_sha256" in manifest
    assert sorted(p.name for p in tmp_path.iterdir()) == ["OUTPUT_FINAL_ROUND.csv"]


def test_display_values_follow_the_template(monkeypatch):
    from packages.export import final_round as fr

    # Seed files cut long names at 200 characters; the full name comes back from ESCAP's baseline.
    full = "Notification of the National Cyber Security Committee " + "x" * 160 + " B.E. 2564 (ประกาศ พ.ศ. 2564)"
    monkeypatch.setattr(fr, "_baseline_names", lambda: (full, full + " ; another instrument"))
    assert fr.law_name(full[:200]) == full
    assert fr.law_name("SEBI CSCRF 2024").startswith("Securities and Exchange Board of India")
    assert fr.law_name("Digital Economy Partnership Agreement") == "Digital Economy Partnership Agreement 2020"
    assert fr.law_name("Privacy Act 1988") == "Privacy Act 1988"
    # Last Amended is a year.
    assert fr.last_amended("2026-07", "Privacy Act 1988") == "2026"
    assert fr.last_amended("", "Computer Crime Act B.E. 2550 as amended B.E. 2560 (ETDA consolidated)") == "2017"
    assert fr.last_amended("B.E. 2562", "x") == "2019" and fr.last_amended(None, "x") == ""
    # Notes in plain words, no internal identifiers.
    assert fr.plain_notes("Discovery: instrument is known but s. sch1cl41 is not among its recorded provisions "
                          "(['10', '110']). Modality: must; exceptions: none") == (
        "New provision: the 2025 RDTII database has this law but not this provision. Operative wording: must.")
    assert fr.plain_notes("Discovery: master dataset already records s. 25A of this law. Modality: n/a; "
                          "exceptions: none") == "Provision already in the 2025 RDTII database."
    assert fr.plain_notes("Discovery: instrument not in the master dataset for India. Modality: shall; exceptions: "
                          "Group B bids are ineligible.").endswith("Exception: Group B bids are ineligible.")
    assert fr.reviewer_text("Genuine safeguard. [Final joint-session workbook, 20 Jul — supersedes any prior "
                            "batch entry.]") == "Genuine safeguard."
    assert fr.model_route("openai/gpt-5.6-luna/escalate:openai/gpt-5.6-terra+text-embedding-3-small") == (
        "gpt-5.6-luna; escalation: gpt-5.6-terra; embeddings: text-embedding-3-small")
    assert fr._gate_summary([{"gate_id": "G1", "status": "PASS"}, {"gate_id": "G6", "status": "WARN"}]) == (
        "1 of 2 checks passed; reviewer judged: must/may wording")
